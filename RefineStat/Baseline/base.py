#!/usr/bin/env python3
"""
base.py — Full baseline runner with execution, diagnostics, and aggregated reporting.

Generates PyMC model code using Hugging Face causal LM sampling (unconstrained), executes the code,
computes Bayesian diagnostics (R-hat, ESS, divergences, BFMI, Pareto-k, ELPD), and produces
aggregated summaries. Serves as a comprehensive baseline for statistical quality assessment.

Key features:
  - Single-pass code generation per dataset
  - In-process execution with full diagnostic extraction
  - Per-dataset results with conditional formatting (Excel)
  - Aggregated reliability and token usage across seeds
  - Token budget calculation (20% buffer from max observed)
  - Leaderboard and progression tracking

Usage:
    python Baseline/base.py --seeds 10 --temperature 0.3
    python Baseline/base.py --seeds 5 --temperature 0.2 --output results/baseline
    python Baseline/base.py --seeds 3 --models "meta-llama/Meta-Llama-3-8B,google/codegemma-7b"
"""
import os
import sys
import re
import json
import gc
import argparse
from typing import Dict, Tuple, Union, List, Any, Optional

import numpy as np
import pandas as pd
import xarray as xr
import arviz as az
import torch
import tiktoken
from openpyxl.styles import PatternFill
from openpyxl import load_workbook
from transformers import AutoModelForCausalLM, AutoTokenizer

parent_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if parent_dir not in sys.path:
    sys.path.append(parent_dir)

from commons.data_pymc import datas_info, build_prompt_generic
from commons.config import config
from commons.utils import (
    convert_np_types, set_seed, check_model_reliability, run_pymc_code
)


# --- Hugging Face sampling (self-contained) ---
_HF_CACHE = os.environ.get("HF_CACHE", "cache/")
_HF_ACCESS_TOKEN = os.environ.get("HF_ACCESS_TOKEN")
_HF_LOAD_KW = dict(
    cache_dir=_HF_CACHE, token=_HF_ACCESS_TOKEN, trust_remote_code=True
)
_BASELINE_LM_IGNORE = frozenset(
    {
        "mode",
        "grammar",
        "chat_mode",
        "parse_output_only",
        "dev_mode",
        "log_level",
        "new_mask_store",
        "parser",
    }
)


class BaselineLM:
    def __init__(
        self,
        model: str,
        *,
        quantize: bool = True,
        device: str = "cuda",
        num_samples: int = 1,
        **gen_kwargs,
    ):
        self._model = (
            AutoModelForCausalLM.from_pretrained(
                model, torch_dtype=torch.bfloat16, **_HF_LOAD_KW
            )
            .eval()
            .to(device)
            if quantize
            else AutoModelForCausalLM.from_pretrained(model, **_HF_LOAD_KW)
            .eval()
            .to(device)
        )
        self.tokenizer = AutoTokenizer.from_pretrained(model, **_HF_LOAD_KW)
        if self.tokenizer.pad_token_id is None and self.tokenizer.eos_token_id is not None:
            self.tokenizer.pad_token = self.tokenizer.eos_token
        self.device = device
        self.num_samples = num_samples
        g = {k: v for k, v in gen_kwargs.items() if k not in _BASELINE_LM_IGNORE}
        g.setdefault("max_new_tokens", 200)
        self._gen_kwargs = g

    @torch.inference_mode()
    def infer(self, prompt):
        if prompt is None:
            raise ValueError("prompt is required")
        batch = [prompt for _ in range(self.num_samples)]
        inputs = self.tokenizer(batch, return_tensors="pt", padding=True).to(self.device)
        cutoff = inputs.input_ids.shape[1]
        gen_kw = dict(self._gen_kwargs)
        if self.tokenizer.pad_token_id is not None:
            gen_kw.setdefault("pad_token_id", self.tokenizer.pad_token_id)
        ids = self._model.generate(**inputs, **gen_kw)
        return [
            self.tokenizer.decode(ids[i, cutoff:], skip_special_tokens=True)
            for i in range(self.num_samples)
        ]


# Default models
DEFAULT_MODELS: List[str] = [
    "meta-llama/Meta-Llama-3-8B",
    "google/codegemma-7b",
    "Qwen/Qwen2.5-Coder-7B",
    "deepseek-ai/DeepSeek-R1-Distill-Qwen-7B",
]

# Diagnostic thresholds for conditional formatting
THRESHOLDS = {
    "r_hat": 1.05,
    "ess_bulk": 400,
    "ess_tail": 100,
    "divergences": 0,
    "bfmi": 0.3,
    "pareto_k": 0.2,
}

RED_FILL = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")


# ---------------------------------------------------------------------------
# Utility Helpers
# ---------------------------------------------------------------------------

def _extract_xarray_value(obj: Any) -> Any:
    """Extract the actual value from an xarray.DataArray if needed."""
    if isinstance(obj, xr.DataArray):
        if obj.size == 1:
            return float(obj.values)
        return obj.values.tolist()
    return obj


def _get_tokenizer(model_name: str):
    """
    Get tokenizer for model, trying tiktoken first, then huggingface.
    
    Args:
        model_name: Model identifier (e.g., "meta-llama/Meta-Llama-3-8B")
    
    Returns:
        Tokenizer instance and count function
    """
    try:
        enc = tiktoken.encoding_for_model(model_name)
        return enc, lambda x: len(enc.encode(x))
    except Exception:
        tokenizer = AutoTokenizer.from_pretrained(model_name, use_fast=True)
        return tokenizer, lambda x: len(tokenizer(x, return_tensors="pt").input_ids[0])


def get_next_experiment_folder(base_dir: str) -> str:
    """Return the next auto-numbered experiment directory."""
    os.makedirs(base_dir, exist_ok=True)
    existing = [
        int(d) for d in os.listdir(base_dir)
        if d.isdigit() and os.path.isdir(os.path.join(base_dir, d))
    ]
    folder = os.path.join(base_dir, str(max(existing, default=0) + 1))
    os.makedirs(folder)
    return folder


def _save_json(data: Any, path: str, label: str = ""):
    """Helper to save JSON with consistent formatting."""
    with open(path, "w") as f:
        json.dump(convert_np_types(data), f, indent=2)
    if label:
        print(f"  Saved {label}")


def _store_token_count(seed: int, dataset: str, llm_key: str, tokens: int, agg: Dict):
    """Store token count in aggregation dictionary with cumulative tracking."""
    agg.setdefault(dataset, {})
    agg[dataset].setdefault(llm_key, {"seeds": {}, "cumulative": {}})
    agg[dataset][llm_key]["seeds"][seed] = tokens
    
    previous = max(
        (agg[dataset][llm_key]["cumulative"].get(s, 0)
         for s in agg[dataset][llm_key]["cumulative"] if s < seed),
        default=0
    )
    agg[dataset][llm_key]["cumulative"][seed] = previous + tokens


def insert_model_code(
    template: str, raw_snippet: Union[str, list], trace_var: str, model_name: str
) -> Tuple[str, int]:
    """
    Insert generated code snippet into PyMC template and count tokens.
    
    Args:
        template: PyMC boilerplate containing 'with pm.Model() as m:'
        raw_snippet: Generated code snippet (string or list)
        trace_var: Variable name for pm.sample result
        model_name: Model ID for tokenizer selection
    
    Returns:
        Tuple of (full_code, token_count)
    
    Raises:
        ValueError: If template marker not found
    """
    # Convert list to string
    if isinstance(raw_snippet, list):
        raw_snippet = "".join(raw_snippet)
    
    # Process lines with indentation
    lines = raw_snippet.replace("```", "").splitlines()
    processed = [
        "" if not L.strip() else "\t" + L.lstrip()
        for L in lines
    ]
    snippet_code = "\n".join(processed)
    
    # Count tokens
    _, count_fn = _get_tokenizer(model_name)
    tok_count = count_fn(snippet_code)
    
    # Insert into template
    out_lines, inserted = [], False
    for bl in template.splitlines():
        out_lines.append(bl.lstrip())
        if bl.strip().startswith("with pm.Model() as m:"):
            out_lines.append(snippet_code)
            inserted = True
    
    if not inserted:
        raise ValueError("'with pm.Model() as m:' not found in template")
    
    out_lines.append(f"\tsummary = az.summary({trace_var})")
    return "\n".join(out_lines), tok_count


def _flatten_diagnostics(diagnostics: Dict[str, Any]) -> Dict[str, Any]:
    """Flatten nested diagnostics dictionary for DataFrame storage."""
    result = {}
    
    # Top-level metrics
    for key in [
        "max_r_hat", "min_ess_bulk", "min_ess_tail", "n_divergent",
        "prop_high_pareto_k", "elpd_loo", "loo_se", "reliability_score"
    ]:
        if key in diagnostics:
            result[key] = _extract_xarray_value(diagnostics[key])
    
    # BFMI values (scalar or array)
    if "bfmi_values" in diagnostics:
        bfmi_values = diagnostics["bfmi_values"]
        if isinstance(bfmi_values, (list, np.ndarray)):
            if isinstance(bfmi_values, np.ndarray):
                bfmi_values = bfmi_values.tolist()
            result["min_bfmi"] = min(bfmi_values)
            result["max_bfmi"] = max(bfmi_values)
            for i, val in enumerate(bfmi_values):
                result[f"bfmi_{i+1}"] = val
    
    return result


def _apply_conditional_formatting(df: pd.DataFrame, filepath: str):
    """Apply red highlight formatting to failing metrics in Excel."""
    # Save base DataFrame
    df.to_excel(filepath, index=False)
    
    # Open for formatting
    wb = load_workbook(filepath)
    ws = wb.active
    
    # Get column indices
    col_indices = {col: i + 1 for i, col in enumerate(df.columns)}
    
    # Apply highlighting per row
    for row in range(2, len(df) + 2):
        row_data = df.iloc[row - 2]
        if pd.isnull(row_data.get("reliability_score")):
            continue
        
        # R-hat check
        if "max_r_hat" in col_indices and pd.notnull(row_data.get("max_r_hat")):
            if row_data["max_r_hat"] >= THRESHOLDS["r_hat"]:
                ws.cell(row=row, column=col_indices["max_r_hat"]).fill = RED_FILL
        
        # ESS bulk check
        if "min_ess_bulk" in col_indices and pd.notnull(row_data.get("min_ess_bulk")):
            if row_data["min_ess_bulk"] < THRESHOLDS["ess_bulk"]:
                ws.cell(row=row, column=col_indices["min_ess_bulk"]).fill = RED_FILL
        
        # ESS tail check
        if "min_ess_tail" in col_indices and pd.notnull(row_data.get("min_ess_tail")):
            if row_data["min_ess_tail"] < THRESHOLDS["ess_tail"]:
                ws.cell(row=row, column=col_indices["min_ess_tail"]).fill = RED_FILL
        
        # Divergences check
        if "n_divergent" in col_indices and pd.notnull(row_data.get("n_divergent")):
            if row_data["n_divergent"] > THRESHOLDS["divergences"]:
                ws.cell(row=row, column=col_indices["n_divergent"]).fill = RED_FILL
        
        # BFMI checks
        for bfmi_col in [c for c in df.columns if c.startswith("bfmi_")]:
            if bfmi_col in col_indices and pd.notnull(row_data.get(bfmi_col)):
                if row_data[bfmi_col] <= THRESHOLDS["bfmi"]:
                    ws.cell(row=row, column=col_indices[bfmi_col]).fill = RED_FILL
        
        # Pareto-k check
        if "prop_high_pareto_k" in col_indices and pd.notnull(row_data.get("prop_high_pareto_k")):
            if row_data["prop_high_pareto_k"] > THRESHOLDS["pareto_k"]:
                ws.cell(row=row, column=col_indices["prop_high_pareto_k"]).fill = RED_FILL
    
    wb.save(filepath)


# ---------------------------------------------------------------------------
# Main Experiment
# ---------------------------------------------------------------------------

def _run_single_seed(
    root: str,
    llm_key: str,
    llm_name: str,
    baseline_model: BaselineLM,
    seed: int,
    reliability_agg: Dict,
    token_agg: Dict,
) -> Dict[str, Dict]:
    """
    Run single seed across all datasets for one model.
    
    Returns:
        Dictionary mapping dataset_name → results_dict
    """
    set_seed(seed)
    exp_dir = os.path.join(root, f"seed_{seed}", llm_key)
    os.makedirs(exp_dir, exist_ok=True)
    
    seed_results = {}
    
    for entry in datas_info:
        dataset = entry["name"]
        ds_dir = os.path.join(exp_dir, dataset)
        os.makedirs(ds_dir, exist_ok=True)
        
        # Initialize result record
        seed_results[dataset] = {
            "seed": seed,
            "model": llm_key,
            "dataset": dataset,
            "total_tokens": 0,
            "cumulative_tokens": 0,
            "code_compiles": False,
            "reliability_score": None,
        }
        
        try:
            # Build prompt
            prompt = build_prompt_generic(entry)
            with open(os.path.join(ds_dir, "prompt.txt"), "w") as f:
                f.write(prompt)
            
            # Generate snippet
            snippet = baseline_model.infer(prompt)
            with open(os.path.join(ds_dir, "snippet.txt"), "w") as f:
                f.write(str(snippet))
            
            # Extract trace variable name
            m = re.search(r"(\w+)\s*=\s*pm.sample", str(snippet))
            trace_var = m.group(1) if m else "trace"
            
            # Insert into template and count tokens
            full_code, tk_count = insert_model_code(
                entry["template_code"], snippet, trace_var, llm_name
            )
            with open(os.path.join(ds_dir, "final_code.py"), "w") as f:
                f.write(full_code)
            
            seed_results[dataset]["total_tokens"] = tk_count
            
            # Execute code
            compiled, out = run_pymc_code(full_code)
            seed_results[dataset]["code_compiles"] = compiled
            
            # Extract diagnostics if successful
            if compiled and isinstance(out, dict) and trace_var in out:
                rel, diag = check_model_reliability(out[trace_var])
                
                # Store in aggregates
                reliability_agg.setdefault(dataset, {})\
                    .setdefault(llm_key, {})[seed] = {
                        "reliability_score": rel,
                        "elpd": _extract_xarray_value(diag.get("elpd_loo")),
                    }
                
                # Update result record
                seed_results[dataset]["reliability_score"] = rel
                
                # Flatten and add diagnostics
                flat_diag = _flatten_diagnostics(diag)
                seed_results[dataset].update(flat_diag)
                
                # Save diagnostics
                diag_to_save = convert_np_types(diag)
                diag_to_save["reliability_score"] = rel
                with open(os.path.join(ds_dir, "diagnostics.json"), "w") as f:
                    json.dump(diag_to_save, f, indent=2)
            
            # Store token count
            _store_token_count(seed, dataset, llm_key, tk_count, token_agg)
            
        except Exception as e:
            error_msg = f"Error: {str(e)}"
            print(f"    {dataset:<20} ✗ {error_msg[:50]}")
            ds_dir = os.path.join(exp_dir, dataset)
            os.makedirs(ds_dir, exist_ok=True)
            with open(os.path.join(ds_dir, "error.txt"), "w") as f:
                f.write(error_msg)
    
    # Update cumulative tokens for each dataset
    for dataset in seed_results:
        cumulative = (
            token_agg.get(dataset, {})
            .get(llm_key, {})
            .get("cumulative", {})
            .get(seed, 0)
        )
        seed_results[dataset]["cumulative_tokens"] = cumulative
    
    # Save per-dataset results (append mode)
    for dataset, results in seed_results.items():
        ds_results_dir = os.path.join(root, llm_key, dataset)
        os.makedirs(ds_results_dir, exist_ok=True)
        
        results_path = os.path.join(ds_results_dir, "results.csv")
        
        # Load existing or create new DataFrame
        if os.path.exists(results_path):
            try:
                df = pd.read_csv(results_path)
                if seed in df["seed"].values:
                    df.loc[df["seed"] == seed] = pd.Series(results)
                else:
                    df = pd.concat([df, pd.DataFrame([results])], ignore_index=True)
            except Exception:
                df = pd.DataFrame([results])
        else:
            df = pd.DataFrame([results])
        
        # Ensure column ordering
        essential_cols = [
            "seed", "reliability_score", "total_tokens", "cumulative_tokens",
            "code_compiles", "elpd_loo"
        ]
        for col in essential_cols:
            if col not in df.columns:
                df[col] = None
        
        remaining_cols = [c for c in df.columns if c not in essential_cols]
        df = df[essential_cols + sorted(remaining_cols)]
        df = df.sort_values("seed")
        
        # Save CSV and Excel
        df.to_csv(results_path, index=False)
        excel_path = os.path.join(ds_results_dir, "results.xlsx")
        _apply_conditional_formatting(df, excel_path)
    
    return seed_results


# ---------------------------------------------------------------------------
# Main Pipeline
# ---------------------------------------------------------------------------

def run_experiment(
    models: List[str],
    temperatures: List[float],
    num_seeds: int,
    output_dir: str,
):
    """
    Run full baseline experiment across models, datasets, and seeds.
    
    Args:
        models: List of model IDs to test
        temperatures: List of temperatures to test
        num_seeds: Number of seeds to run
        output_dir: Base output directory
    """
    num_datasets = len(datas_info)
    
    for temperature in temperatures:
        print(f"\n{'═' * 70}")
        print("  HF Baseline  —  Full Diagnostic Analysis")
        print(f"{'═' * 70}")
        print(f"  Temperature     : {temperature}")
        print(f"  Models          : {len(models)}")
        print(f"  Seeds per model : {num_seeds}")
        print(f"  Datasets        : {num_datasets}  ({', '.join(d['name'] for d in datas_info)})")
        print(f"{'═' * 70}\n")
        
        experiment_dir = get_next_experiment_folder(output_dir)
        
        # Save config
        cfg = {
            "temperature": temperature,
            "models": models,
            "num_seeds": num_seeds,
            "datasets": [d["name"] for d in datas_info],
        }
        with open(os.path.join(experiment_dir, "config.json"), "w") as f:
            json.dump(cfg, f, indent=2)
        
        # Aggregates
        reliability_agg = {}
        token_agg = {}
        
        for llm_id in models:
            print(f"\n{'█' * 70}")
            print(f"  MODEL: {llm_id}")
            print(f"{'█' * 70}")
            
            llm_key = llm_id.replace("/", "_")
            
            try:
                # Load model
                baseline_model = BaselineLM(
                    model=llm_id,
                    do_sample=True,
                    temperature=temperature,
                    device="cuda",
                    max_new_tokens=config.get("max_tokens", 400),
                )
                
                for seed in range(1, num_seeds + 1):
                    print(f"\n  [Seed {seed}/{num_seeds}]")
                    
                    try:
                        _run_single_seed(
                            experiment_dir, llm_key, llm_id, baseline_model,
                            seed, reliability_agg, token_agg,
                        )
                        
                        # Print seed summary
                        datasets_success = 0
                        for data in datas_info:
                            ds_dir = os.path.join(
                                experiment_dir, f"seed_{seed}", llm_key, data["name"]
                            )
                            if os.path.exists(os.path.join(ds_dir, "diagnostics.json")):
                                datasets_success += 1
                        print(f"    ✓ {datasets_success}/{num_datasets} datasets compiled")
                    
                    except Exception as e:
                        print(f"  [Seed {seed}] ERROR: {e}")
                
                # Cleanup
                del baseline_model
                torch.cuda.empty_cache()
                gc.collect()
            
            except Exception as e:
                print(f"  ERROR loading {llm_id}: {e}")
        
        # Save aggregates
        print(f"\n{'─' * 70}")
        print("  Saving aggregates...")
        print(f"{'─' * 70}")
        
        _save_json(reliability_agg, os.path.join(experiment_dir, "aggregated_reliability.json"),
                  "aggregated_reliability.json")
        _save_json(token_agg, os.path.join(experiment_dir, "aggregated_token_count.json"),
                  "aggregated_token_count.json")
        
        # Generate token budget (20% buffer)
        token_budget = {}
        for dataset, dataset_models in token_agg.items():
            for model, data in dataset_models.items():
                if "cumulative" in data:
                    if model not in token_budget:
                        token_budget[model] = {}
                    max_tokens = max(data["cumulative"].values()) if data["cumulative"] else 0
                    token_budget[model][dataset] = int(max_tokens * 1.2)
        
        _save_json(token_budget, os.path.join(experiment_dir, "token_budget.json"),
                  "token_budget.json")
        
        # Process leaderboards and progressions
        datasets = [e["name"] for e in datas_info]
        for llm_id in models:
            llm_key = llm_id.replace("/", "_")
            llm_folder = os.path.join(experiment_dir, llm_key)
            os.makedirs(llm_folder, exist_ok=True)
            
            # Leaderboard
            leaderboard = []
            for s in range(1, num_seeds + 1):
                for ds in datasets:
                    met = reliability_agg.get(ds, {}).get(llm_key, {}).get(s)
                    rel = met["reliability_score"] if met else None
                    elpd = met["elpd"] if met else None
                    cum = (
                        token_agg.get(ds, {})
                        .get(llm_key, {})
                        .get("cumulative", {})
                        .get(s)
                    )
                    code_path = os.path.join(
                        experiment_dir, f"seed_{s}", llm_key, ds, "final_code.py"
                    )
                    if not os.path.exists(code_path):
                        code_path = None
                    
                    leaderboard.append({
                        "seed": s,
                        "dataset": ds,
                        "reliability_score": rel,
                        "elpd_loo": elpd,
                        "cumulative_tokens": cum,
                        "code_path": code_path,
                    })
            
            leaderboard.sort(
                key=lambda x: (
                    x["reliability_score"] or -1,
                    x["elpd_loo"] or -float("inf"),
                ),
                reverse=True,
            )
            
            with open(os.path.join(llm_folder, "leaderboard.json"), "w") as f:
                json.dump(leaderboard, f, indent=2)
            
            # Progression per dataset
            for ds in datasets:
                cum_map = (
                    token_agg.get(ds, {})
                    .get(llm_key, {})
                    .get("cumulative", {})
                )
                progression = []
                for s in sorted(cum_map):
                    rel = (
                        reliability_agg.get(ds, {})
                        .get(llm_key, {})
                        .get(s, {})
                        .get("reliability_score")
                    )
                    progression.append({
                        "seed": s,
                        "cumulative_tokens": cum_map[s],
                        "reliability_score": rel,
                    })
                
                with open(os.path.join(llm_folder, f"progress_{ds}.json"), "w") as f:
                    json.dump(progression, f, indent=2)
        
        # Print final statistics
        summary_lines = []
        summary_lines.append(f"\n{'═' * 70}")
        summary_lines.append(f"  Temperature {temperature} Results")
        summary_lines.append(f"{'═' * 70}")
        total_runs = len(models) * num_seeds * num_datasets
        successful = sum(
            1 for ds in datas_info
            for model in models
            for s in range(1, num_seeds + 1)
            if reliability_agg.get(ds["name"], {}).get(model.replace("/", "_"), {}).get(s)
        )
        summary_lines.append(f"  Total runs      : {total_runs}")
        summary_lines.append(f"  Successful      : {successful}")
        summary_lines.append(f"  Failed          : {total_runs - successful}")
        summary_lines.append(f"  Success rate    : {successful / total_runs * 100:.1f}%")
        summary_lines.append(f"  Output          : {os.path.relpath(experiment_dir)}")
        summary_lines.append(f"{'═' * 70}\n")
        
        summary_text = "\n".join(summary_lines)
        print(summary_text)
        
        # Save summary to file
        summary_file = os.path.join(experiment_dir, "summary.txt")
        with open(summary_file, "w") as f:
            f.write(summary_text)
        print(f"  Saved summary: {summary_file}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Full baseline with execution, diagnostics, and aggregated reporting."
    )
    parser.add_argument(
        "--seeds", "-s", type=int, default=10,
        help="Number of seeds to run per model."
    )
    parser.add_argument(
        "--temperature", "-t", type=float, nargs="+", default=[0.3],
        help="Temperature(s) to test (can specify multiple)."
    )
    parser.add_argument(
        "--output", "-o", type=str, default="results/baseline",
        help="Base output directory."
    )
    parser.add_argument(
        "--models", "-m", type=str, default=None,
        help="Comma-separated model IDs. Defaults to DEFAULT_MODELS."
    )
    
    args = parser.parse_args()
    
    models = (
        [m.strip() for m in args.models.split(",")]
        if args.models
        else DEFAULT_MODELS
    )
    
    run_experiment(
        models=models,
        temperatures=args.temperature,
        num_seeds=args.seeds,
        output_dir=args.output,
    )
