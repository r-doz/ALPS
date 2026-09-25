#!/usr/bin/env python3
"""
base-multi.py — Multi-run HF baseline with best-of-K selection and cross-run aggregation.

Runs Hugging Face unconstrained generation across N independent runs, each with K seeds.
Within each run, selects the best program per dataset (highest reliability_score, elpd_loo
as tiebreaker). After all runs, computes mean ± std across runs to produce results directly
comparable to refinestat/main.py + aggregate_stats.py.

Architecture:
  - N runs × K seeds per run  (seeds are non-overlapping: run 1→seeds 1-K, run 2→K+1-2K, etc.)
  - Within each run: best program per dataset across K seeds
  - Across runs: mean ± std of best-per-run metrics

Outputs (under results/Baseline/base-multi/<expt_N>/):
  - run_<i>/seed_<s>/<model>/<dataset>/  — per-seed generated code + diagnostics
  - run_<i>/<model>/<dataset>/best_program.py  — best program within this run
  - analysis/all_runs_best.csv (.xlsx)  — one row per run × dataset (best from each run)
  - analysis/global_best.csv (.xlsx)  — overall best per dataset across all runs
  - analysis/aggregated_stats.xlsx  — mean ± std across runs (Summary + Formatted sheets)
  - summary.txt

Usage:
    python Baseline/base-multi.py --runs 5 --seeds-per-run 5
    python Baseline/base-multi.py --runs 5 --seeds-per-run 5 --temperature 0.3
    python Baseline/base-multi.py --runs 3 --seeds-per-run 3 --models "Qwen/Qwen2.5-Coder-7B"
"""
import os
import sys
import re
import json
import gc
import argparse
from typing import Dict, List, Any, Optional, Tuple

import numpy as np
import pandas as pd
import xarray as xr
import torch
import tiktoken
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from transformers import AutoModelForCausalLM, AutoTokenizer

parent_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if parent_dir not in sys.path:
    sys.path.append(parent_dir)

from commons.data_pymc import datas_info, build_prompt_generic
from commons.config import config
from commons.utils import convert_np_types, set_seed, check_model_reliability, run_pymc_code


# --- Hugging Face sampling (self-contained) ---
_HF_CACHE = os.environ.get("HF_CACHE", "cache/")
_HF_ACCESS_TOKEN = os.environ.get("HF_ACCESS_TOKEN")
_HF_LOAD_KW = dict(
    cache_dir=_HF_CACHE, token=_HF_ACCESS_TOKEN, trust_remote_code=True
)
_BASELINE_LM_IGNORE = frozenset(
    {
        "mode",
        "grammar",
        "chat_mode",
        "parse_output_only",
        "dev_mode",
        "log_level",
        "new_mask_store",
        "parser",
    }
)


class BaselineLM:
    def __init__(
        self,
        model: str,
        *,
        quantize: bool = True,
        device: str = "cuda",
        num_samples: int = 1,
        **gen_kwargs,
    ):
        self._model = (
            AutoModelForCausalLM.from_pretrained(
                model, torch_dtype=torch.bfloat16, **_HF_LOAD_KW
            )
            .eval()
            .to(device)
            if quantize
            else AutoModelForCausalLM.from_pretrained(model, **_HF_LOAD_KW)
            .eval()
            .to(device)
        )
        self.tokenizer = AutoTokenizer.from_pretrained(model, **_HF_LOAD_KW)
        if self.tokenizer.pad_token_id is None and self.tokenizer.eos_token_id is not None:
            self.tokenizer.pad_token = self.tokenizer.eos_token
        self.device = device
        self.num_samples = num_samples
        g = {k: v for k, v in gen_kwargs.items() if k not in _BASELINE_LM_IGNORE}
        g.setdefault("max_new_tokens", 200)
        self._gen_kwargs = g

    @torch.inference_mode()
    def infer(self, prompt):
        if prompt is None:
            raise ValueError("prompt is required")
        batch = [prompt for _ in range(self.num_samples)]
        inputs = self.tokenizer(batch, return_tensors="pt", padding=True).to(self.device)
        cutoff = inputs.input_ids.shape[1]
        gen_kw = dict(self._gen_kwargs)
        if self.tokenizer.pad_token_id is not None:
            gen_kw.setdefault("pad_token_id", self.tokenizer.pad_token_id)
        ids = self._model.generate(**inputs, **gen_kw)
        return [
            self.tokenizer.decode(ids[i, cutoff:], skip_special_tokens=True)
            for i in range(self.num_samples)
        ]


DEFAULT_MODELS: List[str] = [
    "meta-llama/Meta-Llama-3-8B",
]

THRESHOLDS = {
    "r_hat": 1.05,
    "ess_bulk": 400,
    "ess_tail": 100,
    "divergences": 0,
    "bfmi": 0.3,
    "pareto_k": 0.2,
}

RED_FILL = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")

PRIORITY_METRICS = [
    "reliability_score", "elpd_loo", "max_r_hat", "min_ess_bulk",
    "min_ess_tail", "n_divergent", "prop_high_pareto_k",
    "min_bfmi", "max_bfmi", "total_tokens",
]

AGG_THRESHOLDS = {
    "max_r_hat":          (">=", 1.05),
    "min_ess_bulk":       ("<",  400),
    "min_ess_tail":       ("<",  100),
    "n_divergent":        (">",  0),
    "prop_high_pareto_k": (">",  0.2),
    "min_bfmi":           ("<=", 0.3),
}

FILL_MEAN = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
FILL_STD  = PatternFill(start_color="EBF1DE", end_color="EBF1DE", fill_type="solid")
FILL_META = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
FILL_RED  = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
FILL_HEAD = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


# ---------------------------------------------------------------------------
# Utility helpers
# ---------------------------------------------------------------------------

def _extract_xarray_value(obj: Any) -> Any:
    if isinstance(obj, xr.DataArray):
        return float(obj.values) if obj.size == 1 else obj.values.tolist()
    return obj


def _get_tokenizer(model_name: str):
    try:
        enc = tiktoken.encoding_for_model(model_name)
        return enc, lambda x: len(enc.encode(x))
    except Exception:
        tokenizer = AutoTokenizer.from_pretrained(model_name, use_fast=True)
        return tokenizer, lambda x: len(tokenizer(x, return_tensors="pt").input_ids[0])


def get_next_experiment_folder(base_dir: str) -> str:
    os.makedirs(base_dir, exist_ok=True)
    existing = [
        int(d) for d in os.listdir(base_dir)
        if d.isdigit() and os.path.isdir(os.path.join(base_dir, d))
    ]
    folder = os.path.join(base_dir, str(max(existing, default=0) + 1))
    os.makedirs(folder)
    return folder


def insert_model_code(
    template: str, raw_snippet, trace_var: str, model_name: str
) -> Tuple[str, int]:
    if isinstance(raw_snippet, list):
        raw_snippet = "".join(raw_snippet)
    lines = raw_snippet.replace("```", "").splitlines()
    processed = ["" if not L.strip() else "\t" + L.lstrip() for L in lines]
    snippet_code = "\n".join(processed)

    _, count_fn = _get_tokenizer(model_name)
    tok_count = count_fn(snippet_code)

    out_lines, inserted = [], False
    for bl in template.splitlines():
        out_lines.append(bl.lstrip())
        if bl.strip().startswith("with pm.Model() as m:"):
            out_lines.append(snippet_code)
            inserted = True
    if not inserted:
        raise ValueError("'with pm.Model() as m:' not found in template")
    out_lines.append(f"\tsummary = az.summary({trace_var})")
    return "\n".join(out_lines), tok_count


def _flatten_diagnostics(diagnostics: Dict[str, Any]) -> Dict[str, Any]:
    result = {}
    for key in ["max_r_hat", "min_ess_bulk", "min_ess_tail", "n_divergent",
                "prop_high_pareto_k", "elpd_loo", "loo_se", "reliability_score"]:
        if key in diagnostics:
            result[key] = _extract_xarray_value(diagnostics[key])
    if "bfmi_values" in diagnostics:
        bfmi_values = diagnostics["bfmi_values"]
        if isinstance(bfmi_values, (list, np.ndarray)):
            if isinstance(bfmi_values, np.ndarray):
                bfmi_values = bfmi_values.tolist()
            result["min_bfmi"] = min(bfmi_values)
            result["max_bfmi"] = max(bfmi_values)
            for i, val in enumerate(bfmi_values):
                result[f"bfmi_{i+1}"] = val
    return result


def _apply_conditional_formatting(df: pd.DataFrame, filepath: str):
    df.to_excel(filepath, index=False)
    wb = load_workbook(filepath)
    ws = wb.active
    col_indices = {col: i + 1 for i, col in enumerate(df.columns)}

    for row in range(2, len(df) + 2):
        row_data = df.iloc[row - 2]
        if pd.isnull(row_data.get("reliability_score")):
            continue
        if "max_r_hat" in col_indices and pd.notnull(row_data.get("max_r_hat")):
            if row_data["max_r_hat"] >= THRESHOLDS["r_hat"]:
                ws.cell(row=row, column=col_indices["max_r_hat"]).fill = RED_FILL
        if "min_ess_bulk" in col_indices and pd.notnull(row_data.get("min_ess_bulk")):
            if row_data["min_ess_bulk"] < THRESHOLDS["ess_bulk"]:
                ws.cell(row=row, column=col_indices["min_ess_bulk"]).fill = RED_FILL
        if "min_ess_tail" in col_indices and pd.notnull(row_data.get("min_ess_tail")):
            if row_data["min_ess_tail"] < THRESHOLDS["ess_tail"]:
                ws.cell(row=row, column=col_indices["min_ess_tail"]).fill = RED_FILL
        if "n_divergent" in col_indices and pd.notnull(row_data.get("n_divergent")):
            if row_data["n_divergent"] > THRESHOLDS["divergences"]:
                ws.cell(row=row, column=col_indices["n_divergent"]).fill = RED_FILL
        for bfmi_col in [c for c in df.columns if c.startswith("bfmi_")]:
            if bfmi_col in col_indices and pd.notnull(row_data.get(bfmi_col)):
                if row_data[bfmi_col] <= THRESHOLDS["bfmi"]:
                    ws.cell(row=row, column=col_indices[bfmi_col]).fill = RED_FILL
        if "prop_high_pareto_k" in col_indices and pd.notnull(row_data.get("prop_high_pareto_k")):
            if row_data["prop_high_pareto_k"] > THRESHOLDS["pareto_k"]:
                ws.cell(row=row, column=col_indices["prop_high_pareto_k"]).fill = RED_FILL
    wb.save(filepath)


# ---------------------------------------------------------------------------
# Best-program selection
# ---------------------------------------------------------------------------

def select_best(candidates: List[Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    """Select best candidate by reliability_score, elpd_loo as tiebreaker."""
    valid = [
        c for c in candidates
        if c.get("reliability_score") is not None
        and isinstance(c.get("reliability_score"), (int, float))
    ]
    if not valid:
        return None
    return sorted(
        valid,
        key=lambda x: (
            x.get("reliability_score", -1),
            x.get("elpd_loo") if isinstance(x.get("elpd_loo"), (int, float)) else -float("inf"),
        ),
        reverse=True,
    )[0]


# ---------------------------------------------------------------------------
# Per-seed per-dataset code generation + execution
# ---------------------------------------------------------------------------

def run_single_seed_dataset(
    baseline_model: BaselineLM,
    entry: Dict,
    llm_name: str,
    seed: int,
    ds_dir: str,
) -> Dict[str, Any]:
    dataset = entry["name"]
    result = {
        "seed": seed,
        "dataset": dataset,
        "total_tokens": 0,
        "compiled": False,
        "reliability_score": None,
    }

    prompt = build_prompt_generic(entry)
    with open(os.path.join(ds_dir, "prompt.txt"), "w") as f:
        f.write(prompt)

    snippet = baseline_model.infer(prompt)
    with open(os.path.join(ds_dir, "snippet.txt"), "w") as f:
        f.write(str(snippet))

    m = re.search(r"(\w+)\s*=\s*pm.sample", str(snippet))
    trace_var = m.group(1) if m else "trace"

    full_code, tk_count = insert_model_code(
        entry["template_code"], snippet, trace_var, llm_name
    )
    with open(os.path.join(ds_dir, "final_code.py"), "w") as f:
        f.write(full_code)

    result["total_tokens"] = tk_count
    result["program"] = full_code

    compiled, out = run_pymc_code(full_code)
    result["compiled"] = compiled

    if compiled and isinstance(out, dict) and trace_var in out:
        rel, diag = check_model_reliability(out[trace_var])
        result["reliability_score"] = rel
        result.update(_flatten_diagnostics(diag))

        diag_to_save = convert_np_types(diag)
        diag_to_save["reliability_score"] = rel
        with open(os.path.join(ds_dir, "diagnostics.json"), "w") as f:
            json.dump(diag_to_save, f, indent=2)

    return result


# ---------------------------------------------------------------------------
# Aggregation (mean ± std across runs)
# ---------------------------------------------------------------------------

def _fails_threshold(metric: str, value) -> bool:
    if metric not in AGG_THRESHOLDS or value is None:
        return False
    if isinstance(value, float) and np.isnan(value):
        return False
    op, threshold = AGG_THRESHOLDS[metric]
    if   op == ">=": return value >= threshold
    elif op == ">":  return value > threshold
    elif op == "<":  return value < threshold
    elif op == "<=": return value <= threshold
    return False


def compute_aggregated_stats(
    all_runs_df: pd.DataFrame,
) -> Tuple[pd.DataFrame, List[str]]:
    numeric = all_runs_df.select_dtypes(include=[np.number]).columns.tolist()
    skip = {"run", "seed", "iteration"}
    available = [c for c in numeric if c not in skip]
    metrics = [m for m in PRIORITY_METRICS if m in available]
    extras = sorted(c for c in available if c not in metrics)
    metrics = metrics + extras

    rows = []
    for (model, dataset), group in all_runs_df.groupby(["model", "dataset"], sort=True):
        row = {"model": model, "dataset": dataset}
        for metric in metrics:
            if metric not in group.columns:
                row[f"{metric}_mean"] = None
                row[f"{metric}_std"]  = None
                continue
            valid = group[metric].dropna()
            row[f"{metric}_mean"] = round(valid.mean(), 4) if len(valid) > 0 else None
            row[f"{metric}_std"]  = round(valid.std(),  4) if len(valid) > 1 else 0.0
        rows.append(row)

    return pd.DataFrame(rows), metrics


def write_aggregated_excel(
    stats_df: pd.DataFrame, metrics: List[str], output_path: str
):
    stats_df.to_excel(output_path, index=False, sheet_name="_tmp")
    wb = load_workbook(output_path)
    if "_tmp" in wb.sheetnames:
        del wb["_tmp"]

    # ── Sheet 1: Summary (alternating mean / std columns) ──
    ws = wb.create_sheet("Summary")
    headers = ["model", "dataset"]
    for m in metrics:
        headers.append(f"{m}\nmean")
        headers.append(f"{m}\nstd")

    for c_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c_idx)
        cell.value     = h
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = FILL_HEAD
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border    = THIN_BORDER

    for r_idx, (_, row) in enumerate(stats_df.iterrows(), start=2):
        ws.cell(row=r_idx, column=1).value = row["model"]
        ws.cell(row=r_idx, column=2).value = row["dataset"]
        for col_idx in range(1, 3):
            ws.cell(row=r_idx, column=col_idx).fill   = FILL_META
            ws.cell(row=r_idx, column=col_idx).font   = Font(bold=True, size=10)
            ws.cell(row=r_idx, column=col_idx).border = THIN_BORDER

        for m_idx, metric in enumerate(metrics):
            mean_col = 3 + m_idx * 2
            std_col  = mean_col + 1
            mean_val = row.get(f"{metric}_mean")
            std_val  = row.get(f"{metric}_std")

            mean_cell = ws.cell(row=r_idx, column=mean_col)
            mean_cell.value  = mean_val
            mean_cell.fill   = FILL_RED if _fails_threshold(metric, mean_val) else FILL_MEAN
            mean_cell.font   = Font(size=10)
            mean_cell.border = THIN_BORDER
            if isinstance(mean_val, float):
                mean_cell.number_format = "0.0000"

            std_cell = ws.cell(row=r_idx, column=std_col)
            std_cell.value  = std_val
            std_cell.fill   = FILL_STD
            std_cell.font   = Font(size=10, color="555555")
            std_cell.border = THIN_BORDER
            if isinstance(std_val, float):
                std_cell.number_format = "0.0000"

    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18
    for m_idx in range(len(metrics)):
        ws.column_dimensions[get_column_letter(3 + m_idx * 2)].width = 14
        ws.column_dimensions[get_column_letter(4 + m_idx * 2)].width = 12

    # ── Sheet 2: Formatted (mean ± std strings) ──
    ws2 = wb.create_sheet("Formatted (mean ± std)")
    headers2 = ["model", "dataset"] + metrics
    for c_idx, h in enumerate(headers2, start=1):
        cell = ws2.cell(row=1, column=c_idx)
        cell.value     = h
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = FILL_HEAD
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border    = THIN_BORDER

    for r_idx, (_, row) in enumerate(stats_df.iterrows(), start=2):
        ws2.cell(row=r_idx, column=1).value = row["model"]
        ws2.cell(row=r_idx, column=2).value = row["dataset"]
        for col_idx in range(1, 3):
            ws2.cell(row=r_idx, column=col_idx).fill   = FILL_META
            ws2.cell(row=r_idx, column=col_idx).font   = Font(bold=True, size=10)
            ws2.cell(row=r_idx, column=col_idx).border = THIN_BORDER

        for m_idx, metric in enumerate(metrics):
            col_idx  = 3 + m_idx
            mean_val = row.get(f"{metric}_mean")
            std_val  = row.get(f"{metric}_std")

            if mean_val is None or (isinstance(mean_val, float) and np.isnan(mean_val)):
                display = "—"
            elif std_val is None or (isinstance(std_val, float) and np.isnan(std_val)) or std_val == 0:
                display = f"{mean_val:.4f}"
            else:
                display = f"{mean_val:.4f} ± {std_val:.4f}"

            cell = ws2.cell(row=r_idx, column=col_idx)
            cell.value     = display
            cell.fill      = FILL_RED if _fails_threshold(metric, mean_val) else FILL_MEAN
            cell.font      = Font(size=10)
            cell.border    = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")

    ws2.freeze_panes = "A2"
    ws2.column_dimensions["A"].width = 28
    ws2.column_dimensions["B"].width = 18
    for m_idx in range(len(metrics)):
        ws2.column_dimensions[get_column_letter(3 + m_idx)].width = 20

    wb.save(output_path)


# ---------------------------------------------------------------------------
# Main experiment
# ---------------------------------------------------------------------------

def run_multi_experiment(
    n_runs: int,
    seeds_per_run: int,
    models: List[str],
    temperature: float,
    output_dir: str,
):
    experiment_dir = get_next_experiment_folder(output_dir)
    num_datasets = len(datas_info)
    total_seeds = n_runs * seeds_per_run

    print(f"\n{'═' * 70}")
    print("  HF Baseline Multi-Run  —  Best-of-K + Cross-Run Aggregation")
    print(f"{'═' * 70}")
    print(f"  Runs            : {n_runs}")
    print(f"  Seeds per run   : {seeds_per_run}")
    print(f"  Total seeds     : {total_seeds}")
    print(f"  Models          : {', '.join(models)}")
    print(f"  Datasets        : {num_datasets}  ({', '.join(d['name'] for d in datas_info)})")
    print(f"  Temperature     : {temperature}")
    print(f"  Output          : {os.path.relpath(experiment_dir)}")
    print(f"{'═' * 70}\n")

    cfg = {
        "n_runs": n_runs,
        "seeds_per_run": seeds_per_run,
        "total_seeds": total_seeds,
        "models": models,
        "temperature": temperature,
        "datasets": [d["name"] for d in datas_info],
    }
    with open(os.path.join(experiment_dir, "config.json"), "w") as f:
        json.dump(cfg, f, indent=2)

    all_runs_best_rows: List[Dict] = []
    all_candidates: Dict[str, Dict[str, List[Dict]]] = {}

    for model_id in models:
        llm_key = model_id.replace("/", "_")
        all_candidates[llm_key] = {}

        print(f"\n{'█' * 70}")
        print(f"  MODEL: {model_id}")
        print(f"{'█' * 70}")

        try:
            baseline_model = BaselineLM(
                model=model_id,
                do_sample=True,
                temperature=temperature,
                device="cuda",
                max_new_tokens=config.get("max_tokens", 400),
            )

            for run_idx in range(1, n_runs + 1):
                run_seeds = list(range(
                    (run_idx - 1) * seeds_per_run + 1,
                    run_idx * seeds_per_run + 1,
                ))
                run_dir = os.path.join(experiment_dir, f"run_{run_idx}")
                os.makedirs(run_dir, exist_ok=True)

                print(f"\n  ┌─ Run {run_idx}/{n_runs}  (seeds {run_seeds})"
                      f" {'─' * max(1, 45 - len(str(run_seeds)))}")

                run_results: Dict[str, List[Dict]] = {d["name"]: [] for d in datas_info}

                for seed in run_seeds:
                    set_seed(seed)
                    print(f"  │  Seed {seed}:")

                    for entry in datas_info:
                        dataset = entry["name"]
                        ds_dir = os.path.join(
                            run_dir, f"seed_{seed}", llm_key, dataset
                        )
                        os.makedirs(ds_dir, exist_ok=True)

                        try:
                            result = run_single_seed_dataset(
                                baseline_model, entry, model_id, seed, ds_dir
                            )
                            result["model"] = llm_key
                            run_results[dataset].append(result)

                            status = "✓" if result["compiled"] else "✗"
                            rel = result.get("reliability_score")
                            rel_str = f"{rel:.1f}" if isinstance(rel, (int, float)) else "—"
                            print(f"  │    {status} {dataset:<20} "
                                  f"rel={rel_str:<6} tokens={result['total_tokens']:,}")

                        except Exception as e:
                            print(f"  │    ! {dataset:<20} ERROR: {e}")
                            with open(os.path.join(ds_dir, "error.txt"), "w") as f:
                                f.write(str(e))

                # Select best per dataset within this run
                print(f"  │")
                print(f"  │  Best per dataset (run {run_idx}):")

                for entry in datas_info:
                    dataset = entry["name"]
                    all_candidates[llm_key].setdefault(dataset, [])
                    candidates = run_results[dataset]
                    best = select_best(candidates)

                    if best:
                        best_row = {k: v for k, v in best.items() if k != "program"}
                        best_row["run"] = run_idx
                        best_row["model"] = llm_key
                        all_runs_best_rows.append(best_row)

                        all_candidates[llm_key][dataset].append(best)

                        best_dir = os.path.join(run_dir, llm_key, dataset)
                        os.makedirs(best_dir, exist_ok=True)
                        if "program" in best:
                            with open(os.path.join(best_dir, "best_program.py"), "w") as f:
                                f.write(best["program"])

                        print(f"  │    ✓ {dataset:<20} seed={best['seed']} "
                              f"rel={best.get('reliability_score', 'N/A')}")
                    else:
                        print(f"  │    ✗ {dataset:<20} (no successful programs)")

                print(f"  └{'─' * 60}")

            del baseline_model
            torch.cuda.empty_cache()
            gc.collect()

        except Exception as e:
            print(f"  ERROR loading {model_id}: {e}")

    # -------------------------------------------------------------------
    # Save all-runs best CSV (one row per run × dataset)
    # -------------------------------------------------------------------
    analysis_dir = os.path.join(experiment_dir, "analysis")
    os.makedirs(analysis_dir, exist_ok=True)

    essential_cols = [
        "model", "dataset", "run", "seed", "reliability_score",
        "elpd_loo", "total_tokens", "max_r_hat", "min_ess_bulk",
        "min_ess_tail", "n_divergent", "prop_high_pareto_k",
    ]

    if all_runs_best_rows:
        all_runs_df = pd.DataFrame(all_runs_best_rows)
        first_cols = [c for c in essential_cols if c in all_runs_df.columns]
        other_cols = [c for c in all_runs_df.columns
                      if c not in first_cols and c != "program"]
        all_runs_df = all_runs_df[first_cols + sorted(other_cols)]

        csv_path  = os.path.join(analysis_dir, "all_runs_best.csv")
        xlsx_path = os.path.join(analysis_dir, "all_runs_best.xlsx")
        all_runs_df.to_csv(csv_path, index=False)
        _apply_conditional_formatting(all_runs_df, xlsx_path)
        print(f"\n  Saved: all_runs_best.csv  ({len(all_runs_df)} rows)")
    else:
        all_runs_df = pd.DataFrame()
        print("\n  No successful runs to save.")

    # -------------------------------------------------------------------
    # Global best per model/dataset (across ALL runs)
    # -------------------------------------------------------------------
    print(f"\n{'─' * 70}")
    print("  Global best programs (across all runs):")
    print(f"{'─' * 70}")

    global_best_rows = []
    for model_id in models:
        llm_key = model_id.replace("/", "_")
        for entry in datas_info:
            dataset = entry["name"]
            candidates = all_candidates.get(llm_key, {}).get(dataset, [])
            best = select_best(candidates)
            if best:
                save_dir = os.path.join(experiment_dir, llm_key, dataset)
                os.makedirs(save_dir, exist_ok=True)
                program = best.get("program", "")
                if program:
                    with open(os.path.join(save_dir, "best_program.py"), "w") as f:
                        f.write(program)
                best_row = {k: v for k, v in best.items() if k != "program"}
                global_best_rows.append(best_row)
                print(f"  ✓ {llm_key}/{dataset}  "
                      f"(seed {best.get('seed', '?')}, "
                      f"reliability={best.get('reliability_score', 'N/A')})")
            else:
                print(f"  ✗ {llm_key}/{dataset}  (no successful programs)")

    if global_best_rows:
        global_df = pd.DataFrame(global_best_rows)
        first_cols = [c for c in essential_cols if c in global_df.columns]
        other_cols = [c for c in global_df.columns if c not in first_cols]
        global_df = global_df[first_cols + sorted(other_cols)]
        global_df.to_csv(os.path.join(analysis_dir, "global_best.csv"), index=False)
        _apply_conditional_formatting(
            global_df, os.path.join(analysis_dir, "global_best.xlsx")
        )
        print(f"  Saved: global_best.csv  ({len(global_df)} rows)")

    # -------------------------------------------------------------------
    # Aggregate mean ± std across runs
    # -------------------------------------------------------------------
    if not all_runs_df.empty:
        stats_df, metrics = compute_aggregated_stats(all_runs_df)
        agg_path = os.path.join(analysis_dir, "aggregated_stats.xlsx")
        write_aggregated_excel(stats_df, metrics, agg_path)
        print(f"  Saved: aggregated_stats.xlsx")

        # Console summary
        summary_lines = [
            f"\n{'═' * 70}",
            "  Aggregated Statistics (mean ± std across runs)",
            f"{'═' * 70}",
        ]
        display_cols = [
            "dataset", "reliability_score_mean", "elpd_loo_mean",
            "max_r_hat_mean", "min_ess_bulk_mean", "n_divergent_mean",
        ]
        available_cols = [c for c in display_cols if c in stats_df.columns]
        summary_lines.append(stats_df[available_cols].to_string(index=False))
        summary_lines.append(f"{'═' * 70}")

        total_best = len(all_runs_best_rows)
        successful = sum(
            1 for r in all_runs_best_rows
            if r.get("reliability_score") is not None
        )
        summary_lines.extend([
            "",
            f"  Experiment  : HF Baseline Multi-Run",
            f"  Runs        : {n_runs} × {seeds_per_run} seeds/run"
            f" = {total_seeds} total seeds",
            f"  Models      : {', '.join(models)}",
            f"  Temperature : {temperature}",
            f"  Datasets    : {num_datasets}",
            f"  Best found  : {successful}/{total_best} run×dataset combinations",
            f"  Output      : {os.path.relpath(experiment_dir)}",
            f"{'═' * 70}\n",
        ])

        summary_text = "\n".join(summary_lines)
        print(summary_text)

        summary_file = os.path.join(experiment_dir, "summary.txt")
        with open(summary_file, "w") as f:
            f.write(summary_text)

        agg_summary_file = os.path.join(analysis_dir, "aggregated_summary.txt")
        with open(agg_summary_file, "w") as f:
            f.write(summary_text)
    else:
        print("\n  No data to aggregate.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Multi-run HF baseline with best-of-K selection "
                    "and cross-run aggregation."
    )
    parser.add_argument(
        "--runs", "-r", type=int, default=5,
        help="Number of independent runs (default: 5).",
    )
    parser.add_argument(
        "--seeds-per-run", "-k", type=int, default=5,
        help="Seeds per run — best program selected from these (default: 5).",
    )
    parser.add_argument(
        "--temperature", "-t", type=float, default=0.3,
        help="Sampling temperature for generation.",
    )
    parser.add_argument(
        "--output", "-o", type=str, default="results/Baseline/base-multi",
        help="Base output directory.",
    )
    parser.add_argument(
        "--models", "-m", type=str, default=None,
        help="Comma-separated model IDs. Defaults to DEFAULT_MODELS.",
    )
    args = parser.parse_args()

    models = (
        [m.strip() for m in args.models.split(",")]
        if args.models
        else DEFAULT_MODELS
    )

    run_multi_experiment(
        n_runs=args.runs,
        seeds_per_run=args.seeds_per_run,
        models=models,
        temperature=args.temperature,
        output_dir=args.output,
    )
