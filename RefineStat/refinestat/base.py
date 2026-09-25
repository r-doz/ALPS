#!/usr/bin/env python3
"""
base.py — Baseline experiment runner using iterative refinement.

Generates PyMC model code with IterGen + iterative checker feedback (refinegen.base),
assembles each snippet into the dataset boilerplate, executes it in-process, and
computes Bayesian reliability diagnostics. Results are aggregated across seeds and
saved as per-dataset files and a cross-seed statistical summary (CSV/XLSX).

Usage:
    python refinestat/base.py --seeds 10
    python refinestat/base.py --seeds 5 --temperature 0.5 --output results/base
    python refinestat/base.py --seeds 3 --models Qwen/Qwen2.5-Coder-7B
"""
import os
import sys
import json
import gc
import argparse
from typing import Dict, Any, List, Optional

import numpy as np
import torch
import pandas as pd
import xarray as xr
from openpyxl.styles import PatternFill

# Set up paths for imports
current_dir = os.path.dirname(os.path.abspath(__file__))
parent_dir = os.path.dirname(current_dir)

if current_dir not in sys.path:
    sys.path.insert(0, current_dir)
if parent_dir not in sys.path:
    sys.path.append(parent_dir)

from refinegen.base import iterative_refine
from refinegen.itergen.itergen.main import IterGen
from refinegen.checkers.library_specific.ppl_checker import PyMCChecker

from commons.data_pymc import datas_info, build_prompt_generic
from commons.config import config
from commons.utils import convert_np_types, set_seed, check_model_reliability, run_pymc_code


# Default model — override with --models at the command line.
DEFAULT_MODELS: List[str] = [
    "meta-llama/Meta-Llama-3-8B",
    "google/codegemma-7b",
    "Qwen/Qwen2.5-Coder-7B",
    "deepseek-ai/DeepSeek-R1-Distill-Qwen-7B"
]

# Reliability diagnostic thresholds used for Excel conditional formatting.
THRESHOLDS = {
    "r_hat":     1.05,   # max_r_hat must be < 1.05
    "ess_bulk":  400,    # min_ess_bulk must be >= 400
    "ess_tail":  100,    # min_ess_tail must be >= 100
    "divergences": 0,    # n_divergent must be 0
    "bfmi":      0.3,    # bfmi_values must be > 0.3
    "pareto_k":  0.2,    # prop_high_pareto_k must be <= 0.2
}

RED_FILL = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")


# ---------------------------------------------------------------------------
# Utility helpers
# ---------------------------------------------------------------------------

def extract_xarray_value(obj) -> Any:
    """Return a plain Python scalar or list from an xarray.DataArray."""
    if isinstance(obj, xr.DataArray):
        return float(obj.values) if obj.size == 1 else obj.values.tolist()
    return obj


def flatten_diagnostics(diagnostics: Dict[str, Any]) -> Dict[str, Any]:
    """Flatten a nested diagnostics dict into a single-level dict for DataFrame rows."""
    result = {}
    for key in ["max_r_hat", "min_ess_bulk", "min_ess_tail", "n_divergent",
                "prop_high_pareto_k", "elpd_loo", "loo_se", "reliability_score"]:
        if key in diagnostics:
            result[key] = extract_xarray_value(diagnostics[key])
    if "bfmi_values" in diagnostics:
        bfmi = diagnostics["bfmi_values"]
        if isinstance(bfmi, np.ndarray):
            bfmi = bfmi.tolist()
        if isinstance(bfmi, list):
            result["min_bfmi"] = min(bfmi)
            result["max_bfmi"] = max(bfmi)
            for i, val in enumerate(bfmi):
                result[f"bfmi_{i + 1}"] = val
    return result


def create_excel_with_highlighting(df: pd.DataFrame, filepath: str):
    """Save a DataFrame as .xlsx with red fill on cells that fail reliability thresholds."""
    df.to_excel(filepath, index=False)
    from openpyxl import load_workbook
    wb = load_workbook(filepath)
    ws = wb.active
    col_indices = {col: i for i, col in enumerate(df.columns, start=1)}

    threshold_checks = {
        "max_r_hat":          lambda v: v >= THRESHOLDS["r_hat"],
        "min_ess_bulk":       lambda v: v < THRESHOLDS["ess_bulk"],
        "min_ess_tail":       lambda v: v < THRESHOLDS["ess_tail"],
        "n_divergent":        lambda v: v > THRESHOLDS["divergences"],
        "prop_high_pareto_k": lambda v: v > THRESHOLDS["pareto_k"],
    }

    for row_idx in range(2, len(df) + 2):
        row = df.iloc[row_idx - 2]
        if "reliability_score" not in col_indices or pd.isnull(row.get("reliability_score")):
            continue
        for col, fails in threshold_checks.items():
            if col in col_indices and pd.notnull(row.get(col)) and fails(row[col]):
                ws.cell(row=row_idx, column=col_indices[col]).fill = RED_FILL
        for bfmi_col in [c for c in df.columns if c.startswith("bfmi_")]:
            if bfmi_col in col_indices and pd.notnull(row.get(bfmi_col)):
                if row[bfmi_col] <= THRESHOLDS["bfmi"]:
                    ws.cell(row=row_idx, column=col_indices[bfmi_col]).fill = RED_FILL

    wb.save(filepath)


def extract_trace_name(symboltable: dict) -> str:
    """Return the variable name assigned to pymc.sample in the symbol table."""
    for key, value in symboltable.items():
        if value == "pymc.sample":
            return key
    return "trace"


def insert_model_code(boilerplate: str, raw_snippet, trace_name: str) -> str:
    """
    Insert the generated model snippet into the PyMC boilerplate template under
    'with pm.Model() as m:', and append ArviZ diagnostic print statements.
    """
    if isinstance(raw_snippet, list):
        raw_snippet = "".join(raw_snippet)
    model_snippet = raw_snippet.replace("```", "")
    processed_lines = [
        ("" if not line.strip() else "\t" + line.lstrip())
        for line in model_snippet.splitlines()
    ]
    final_snippet = "\n".join(processed_lines)

    new_lines = []
    snippet_inserted = False
    for line in boilerplate.split("\n"):
        new_lines.append(line.lstrip())
        if line.strip().startswith("with pm.Model() as m:"):
            new_lines.append(final_snippet)
            snippet_inserted = True
    if not snippet_inserted:
        raise ValueError("'with pm.Model() as m:' not found in boilerplate template.")

    new_lines.append(
        "\t# Posterior diagnostics\n"
        f"\tprint('R-hat:', az.rhat({trace_name}))\n"
        f"\tprint('ESS:', az.ess({trace_name}))\n"
        f"\tloo = az.loo({trace_name})\n"
        f"\tprint('ELPD LOO:', loo.elpd_loo)\n"
    )
    return "\n".join(new_lines)


# ---------------------------------------------------------------------------
# Aggregation helpers (operate on a shared state dict)
# ---------------------------------------------------------------------------

def _store_reliability(state: dict, seed: int, dataset: str, llm_key: str, data: dict):
    state["reliability"].setdefault(dataset, {}).setdefault(llm_key, {})[seed] = data


def _store_tokens(state: dict, seed: int, dataset: str, llm_key: str, tokens: int):
    entry = state["tokens"].setdefault(dataset, {}).setdefault(
        llm_key, {"seeds": {}, "cumulative": {}}
    )
    entry["seeds"][seed] = tokens
    prev = max(
        (entry["cumulative"].get(s, 0) for s in entry["cumulative"] if s < seed),
        default=0,
    )
    entry["cumulative"][seed] = prev + tokens


def _store_compilation(state: dict, seed: int, dataset: str, llm_key: str,
                       compiled: bool, total_programs: int):
    state["compilation"].setdefault(dataset, {}).setdefault(llm_key, {})[seed] = {
        "compiled": compiled,
        "total_programs": total_programs,
        "compilation_rate": int(compiled) / total_programs if total_programs > 0 else 0.0,
    }


def save_aggregates(output_dir: str, state: dict):
    """Persist the three aggregation dicts as JSON files."""
    mapping = [
        ("aggregated_reliability", "reliability"),
        ("aggregated_token_count", "tokens"),
        ("compilation_stats",      "compilation"),
    ]
    for filename, key in mapping:
        path = os.path.join(output_dir, f"{filename}.json")
        with open(path, "w") as f:
            json.dump(convert_np_types(state[key]), f, indent=2)
        print(f"Saved {path}")


# ---------------------------------------------------------------------------
# Core experiment logic
# ---------------------------------------------------------------------------

def _print_seed_header(seed: int, total_seeds: int, llm_model: str):
    print(f"\n{'─' * 60}")
    print(f"  Seed {seed}/{total_seeds}  |  Model: {llm_model}")
    print(f"{'─' * 60}")


def _print_dataset_row(dataset: str, compiled: bool, reliability, tokens: int):
    status = "✓" if compiled else "✗"
    rel_str = f"{reliability:.1f}" if isinstance(reliability, (int, float)) else "—"
    tok_str = f"{tokens:,}"
    print(f"  {status}  {dataset:<20}  reliability={rel_str:<6}  tokens={tok_str}")


def _print_seed_summary(seed: int, compiled: int, total: int):
    rate = compiled / total * 100 if total else 0
    bar_filled = int(rate / 10)
    bar = "█" * bar_filled + "░" * (10 - bar_filled)
    print(f"\n  [{bar}] {compiled}/{total} datasets compiled ({rate:.0f}%)")


def _write_seed_summary(filepath: str, seed: int, results: list):
    """Write an aligned tabular summary for a single seed to a text file."""
    compiled_results = [r for r in results if r["compiled"]]
    failed_results   = [r for r in results if not r["compiled"]]
    total = len(results)
    compiled_count = len(compiled_results)

    lines = [
        f"Seed: {seed}",
        f"Datasets total    : {total}",
        f"Compiled          : {compiled_count}",
        f"Failed            : {total - compiled_count}",
        f"Compilation rate  : {compiled_count / total * 100:.1f}%",
        "",
        f"{'Dataset':<22} {'Compiled':<10} {'Reliability':>12} {'Tokens':>10} {'ELPD LOO':>12}",
        "─" * 70,
    ]
    for r in results:
        rel = r["reliability_score"]
        elpd = r["diagnostics"].get("elpd_loo") if isinstance(r.get("diagnostics"), dict) else None
        lines.append(
            f"{r['dataset']:<22} "
            f"{'yes' if r['compiled'] else 'no':<10} "
            f"{(f'{rel:.2f}' if isinstance(rel, (int, float)) else 'N/A'):>12} "
            f"{r['total_tokens']:>10,} "
            f"{(f'{float(elpd):.2f}' if elpd is not None else 'N/A'):>12}"
        )
    lines += ["", "Failed datasets:"] + (
        [f"  - {r['dataset']}" for r in failed_results] if failed_results else ["  (none)"]
    )

    with open(filepath, "w") as f:
        f.write("\n".join(lines))


def run_experiment_for_seed(iter_gen, seed: int, total_seeds: int, llm_model: str,
                             parent_folder: str, state: dict) -> dict:
    """
    Run iterative refinement for all datasets for a single seed.
    Updates state in place and returns a per-seed summary dict.
    """
    set_seed(seed)
    experiment_folder = os.path.join(parent_folder, f"seed_{seed}")
    os.makedirs(experiment_folder, exist_ok=True)
    with open(os.path.join(experiment_folder, "seed.txt"), "w") as f:
        f.write(str(seed))

    _print_seed_header(seed, total_seeds, llm_model)
    print(f"  Output: {experiment_folder}\n")

    llm_key = llm_model.replace("/", "_")
    overall_results = []

    for data_entry in datas_info:
        dataset = data_entry["name"]
        pair_folder = os.path.join(experiment_folder, dataset)
        os.makedirs(pair_folder, exist_ok=True)

        prompt = build_prompt_generic(data_entry)
        with open(os.path.join(pair_folder, "prompt.txt"), "w") as f:
            f.write(prompt)

        result = {
            "llm_model": llm_model, "dataset": dataset, "seed": seed,
            "compiled": False, "reliability_score": None,
            "total_tokens": 0, "diagnostics": None,
        }

        try:
            final_code, symboltable, refinement_log = iterative_refine(
                prompt=prompt,
                iter_gen=iter_gen,
                checker_cls=PyMCChecker,
                unit_name="function_call",
                max_iter=config.get("max_iter", 35),
                checker_config={},
                symboltable={"pm": "pymc"},
                dedent=True,
            )

            with open(os.path.join(pair_folder, "refinement_log.txt"), "w") as f:
                f.write(refinement_log)

            trace_name = extract_trace_name(symboltable)
            full_code = insert_model_code(data_entry["template_code"], final_code, trace_name)
            with open(os.path.join(pair_folder, "final_code.py"), "w") as f:
                f.write(full_code)

            compiled, exec_output = run_pymc_code(full_code)
            with open(os.path.join(pair_folder, "exec_output.txt"), "w") as f:
                f.write(str(exec_output))

            total_tokens = iter_gen._metadata.get("total_tokens", 0)
            result["compiled"] = compiled
            result["total_tokens"] = total_tokens
            _store_compilation(state, seed, dataset, llm_key, compiled, 1)
            _store_tokens(state, seed, dataset, llm_key, total_tokens)

            if compiled and isinstance(exec_output, dict) and trace_name in exec_output:
                rel, diag = check_model_reliability(exec_output[trace_name])
                _store_reliability(state, seed, dataset, llm_key, {
                    "reliability_score": rel,
                    "elpd_loo": extract_xarray_value(diag.get("elpd_loo")),
                    "diagnostics": convert_np_types(diag),
                })
                diag_out = convert_np_types(diag)
                diag_out["reliability_score"] = rel
                with open(os.path.join(pair_folder, "diagnostics.json"), "w") as f:
                    json.dump(diag_out, f, indent=2)
                result["reliability_score"] = rel
                result["diagnostics"] = diag

        except Exception as e:
            print(f"  ! ERROR on {dataset}: {e}")
            with open(os.path.join(pair_folder, "error.txt"), "w") as f:
                f.write(str(e))

        overall_results.append(result)
        _print_dataset_row(dataset, result["compiled"],
                           result["reliability_score"], result["total_tokens"])

    compiled_success = sum(1 for r in overall_results if r["compiled"])
    total = len(overall_results)
    _print_seed_summary(seed, compiled_success, total)

    _write_seed_summary(
        os.path.join(experiment_folder, "summary.txt"), seed, overall_results
    )

    return {
        "seed": seed,
        "total": total,
        "compiled_success": compiled_success,
        "compilation_rate": compiled_success / total,
        "experiment_folder": experiment_folder,
        "results": overall_results,
    }


def create_statistical_summary(output_dir: str, state: dict) -> Optional[pd.DataFrame]:
    """Compute cross-seed mean/std statistics and save as CSV and XLSX."""
    reliability_aggregate = state["reliability"]
    token_count_aggregate = state["tokens"]
    compilation_stats = state["compilation"]

    all_results = []
    for dataset, models in reliability_aggregate.items():
        for llm_key, seeds in models.items():
            for seed, data in seeds.items():
                row = {
                    "dataset": dataset, "llm_model": llm_key, "seed": seed,
                    "reliability_score": data["reliability_score"],
                    "elpd_loo": data["elpd_loo"],
                }
                if "diagnostics" in data:
                    row.update(flatten_diagnostics(data["diagnostics"]))
                all_results.append(row)

    df = pd.DataFrame(all_results) if all_results else pd.DataFrame()

    diag_cols = ["max_r_hat", "min_ess_bulk", "min_ess_tail",
                 "n_divergent", "prop_high_pareto_k", "min_bfmi", "max_bfmi"]
    stats_rows = []

    for dataset, models in compilation_stats.items():
        for llm_model, seeds in models.items():
            total_seeds = len(seeds)
            successful_seeds = sum(1 for s in seeds.values() if s["compiled"])
            compilation_rate = successful_seeds / total_seeds if total_seeds > 0 else 0.0

            subset = (
                df[(df["dataset"] == dataset) & (df["llm_model"] == llm_model)]
                if not df.empty else pd.DataFrame()
            )
            token_entry = token_count_aggregate.get(dataset, {}).get(llm_model, {})
            tokens = list(token_entry.get("seeds", {}).values())
            token_mean = np.mean(tokens) if tokens else 0.0
            token_std = np.std(tokens) if tokens else 0.0

            rel_mean = rel_std = elpd_mean = elpd_std = None
            diag_stats = {f"{c}_mean": None for c in diag_cols}
            diag_stats.update({f"{c}_std": None for c in diag_cols})

            if not subset.empty:
                rel_vals = subset["reliability_score"].dropna()
                if not rel_vals.empty:
                    rel_mean, rel_std = rel_vals.mean(), rel_vals.std()
                elpd_vals = subset["elpd_loo"].dropna()
                if not elpd_vals.empty:
                    elpd_mean, elpd_std = elpd_vals.mean(), elpd_vals.std()
                for col in diag_cols:
                    vals = subset[col].dropna() if col in subset.columns else pd.Series(dtype=float)
                    if not vals.empty:
                        diag_stats[f"{col}_mean"] = vals.mean()
                        diag_stats[f"{col}_std"] = vals.std()

            stats_rows.append({
                "dataset": dataset,
                "llm_model": llm_model,
                "total_seeds": total_seeds,
                "successful_seeds": successful_seeds,
                "compilation_rate": compilation_rate,
                "reliability_score_mean": rel_mean,
                "reliability_score_std": rel_std,
                "elpd_loo_mean": elpd_mean,
                "elpd_loo_std": elpd_std,
                "total_tokens_mean": token_mean,
                "total_tokens_std": token_std,
                **diag_stats,
            })

    if not stats_rows:
        print("No compiled results to summarize.")
        return None

    summary_df = pd.DataFrame(stats_rows)
    summary_df.to_csv(os.path.join(output_dir, "statistical_summary.csv"), index=False)
    create_excel_with_highlighting(summary_df, os.path.join(output_dir, "statistical_summary.xlsx"))

    total_programs = sum(
        s["total_programs"]
        for m in compilation_stats.values()
        for seeds in m.values()
        for s in seeds.values()
    )
    total_compiled = sum(
        1
        for m in compilation_stats.values()
        for seeds in m.values()
        for s in seeds.values()
        if s["compiled"]
    )

    print(f"\n{'═' * 70}")
    print("  CROSS-SEED STATISTICAL SUMMARY")
    print(f"{'═' * 70}")
    print(f"  {'Dataset':<22} {'Model':<30} {'Compiled':>8} {'Rel (mean±std)':>16} {'ELPD (mean)':>12}")
    print(f"  {'─' * 66}")
    for row in stats_rows:
        rel_str = (
            f"{row['reliability_score_mean']:.2f}±{row['reliability_score_std']:.2f}"
            if row["reliability_score_mean"] is not None else "N/A"
        )
        elpd_str = (
            f"{row['elpd_loo_mean']:.2f}"
            if row["elpd_loo_mean"] is not None else "N/A"
        )
        compiled_str = f"{row['successful_seeds']}/{row['total_seeds']}"
        print(f"  {row['dataset']:<22} {row['llm_model']:<30} {compiled_str:>8} {rel_str:>16} {elpd_str:>12}")
    print(f"  {'─' * 66}")
    print(f"  Overall: {total_compiled}/{total_programs} programs compiled "
          f"({total_compiled / total_programs * 100:.1f}%)")
    print(f"{'═' * 70}")
    print(f"\n  Saved: {os.path.join(output_dir, 'statistical_summary.csv')}")
    print(f"  Saved: {os.path.join(output_dir, 'statistical_summary.xlsx')}")

    return summary_df


def get_next_experiment_folder(base_dir: str) -> str:
    """Return the next auto-numbered experiment directory under base_dir."""
    os.makedirs(base_dir, exist_ok=True)
    existing = [
        int(d) for d in os.listdir(base_dir)
        if d.isdigit() and os.path.isdir(os.path.join(base_dir, d))
    ]
    folder = os.path.join(base_dir, str(max(existing, default=0) + 1))
    os.makedirs(folder)
    return folder


def run_experiments(models: List[str], num_seeds: int, temperature: float, output_dir: str):
    """Top-level orchestrator: iterate over models and seeds, aggregate, and save summaries."""
    experiment_dir = get_next_experiment_folder(output_dir)

    num_datasets = len(datas_info)
    total_expected = len(models) * num_seeds * num_datasets

    print(f"\n{'═' * 70}")
    print("  RefineStat-Base  —  Iterative Refinement Experiment")
    print(f"{'═' * 70}")
    print(f"  Models          : {', '.join(models)}")
    print(f"  Seeds per model : {num_seeds}")
    print(f"  Datasets        : {num_datasets}  ({', '.join(d['name'] for d in datas_info)})")
    print(f"  Temperature     : {temperature}")
    print(f"  Total runs      : {total_expected}")
    print(f"  Output          : {experiment_dir}")
    print(f"{'═' * 70}\n")

    cfg = {
        "models": models,
        "num_seeds": num_seeds,
        "temperature": temperature,
        "datasets": [d["name"] for d in datas_info],
    }
    with open(os.path.join(experiment_dir, "config.json"), "w") as f:
        json.dump(cfg, f, indent=2)

    state = {"reliability": {}, "tokens": {}, "compilation": {}}
    all_summaries: Dict[str, list] = {m: [] for m in models}

    for llm_model in models:
        print(f"\n{'█' * 70}")
        print(f"  MODEL: {llm_model}")
        print(f"{'█' * 70}")
        model_folder = os.path.join(experiment_dir, llm_model.replace("/", "_"))
        os.makedirs(model_folder, exist_ok=True)

        iter_gen = IterGen(
            grammar=os.path.join(current_dir, "refinegen/itergen/itergen/syncode/syncode/parsers/grammars/python_grammar.lark"),
            model_id=llm_model,
            device="cuda",
            do_sample=True,
            temperature=temperature,
        )

        for seed in range(1, num_seeds + 1):
            summary = run_experiment_for_seed(
                iter_gen, seed, num_seeds, llm_model, model_folder, state
            )
            all_summaries[llm_model].append(summary)

        del iter_gen
        torch.cuda.empty_cache()
        gc.collect()

    print(f"\n{'─' * 70}")
    print("  Saving aggregated results...")
    print(f"{'─' * 70}")
    save_aggregates(experiment_dir, state)
    create_statistical_summary(experiment_dir, state)

    total = sum(s["total"] for summaries in all_summaries.values() for s in summaries)
    compiled = sum(s["compiled_success"] for summaries in all_summaries.values() for s in summaries)

    summary_lines = [
        "RefineStat-Base — Aggregated Experiment Summary",
        "=" * 50,
        f"Models           : {', '.join(models)}",
        f"Seeds per model  : {num_seeds}",
        f"Datasets         : {', '.join(d['name'] for d in datas_info)}",
        f"Temperature      : {temperature}",
        "",
        "Results",
        "-" * 50,
        f"Total runs       : {total}",
        f"Compiled         : {compiled}",
        f"Failed           : {total - compiled}",
        f"Compilation rate : {compiled / total * 100:.1f}%",
        "",
        "Per-model breakdown",
        "-" * 50,
    ]
    for llm_model, summaries in all_summaries.items():
        m_total    = sum(s["total"] for s in summaries)
        m_compiled = sum(s["compiled_success"] for s in summaries)
        summary_lines.append(
            f"  {llm_model:<40} {m_compiled}/{m_total} "
            f"({m_compiled / m_total * 100:.1f}%)"
        )
    summary_lines += ["", f"Output directory : {experiment_dir}"]

    with open(os.path.join(experiment_dir, "aggregated_summary.txt"), "w") as f:
        f.write("\n".join(summary_lines))

    print(f"\n{'═' * 70}")
    print(f"  Experiment complete. All results saved in:")
    print(f"  {experiment_dir}")
    print(f"{'═' * 70}\n")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Baseline iterative-refinement runner for PyMC model synthesis."
    )
    parser.add_argument("--seeds", "-s", type=int, default=10,
                        help="Number of seeds to run per model.")
    parser.add_argument("--temperature", "-t", type=float, default=0.3,
                        help="Sampling temperature for IterGen.")
    parser.add_argument("--output", "-o", type=str, default="results/refinestat-base",
                        help="Base output directory.")
    parser.add_argument("--models", "-m", type=str, default=None,
                        help="Comma-separated model IDs. Defaults to DEFAULT_MODELS.")
    args = parser.parse_args()

    models = [m.strip() for m in args.models.split(",")] if args.models else DEFAULT_MODELS
    run_experiments(
        models=models,
        num_seeds=args.seeds,
        temperature=args.temperature,
        output_dir=args.output,
    )
