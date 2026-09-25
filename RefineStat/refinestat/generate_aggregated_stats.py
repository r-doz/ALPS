#!/usr/bin/env python3
"""
Standalone script to generate aggregated statistics from existing experiment results.

This script reads an existing all_seeds_best_programs_summary.csv file and generates:
1. aggregated_stats.xlsx - Mean and standard deviation across seeds
2. aggregated_summary.txt - Publication-ready text summary

Usage:
    python generate_aggregated_stats.py --input results/refinestat-main/2/analysis/all_seeds_best_programs_summary.csv
"""

import os
import sys
import argparse
import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side


# Columns to aggregate (skip metadata columns)
SKIP_COLS = {"model", "dataset", "seed", "iteration"}

# Core metrics to include in the aggregated output (in this order)
PRIORITY_METRICS = [
    "reliability_score",
    "elpd_loo",
    "max_r_hat",
    "min_ess_bulk",
    "min_ess_tail",
    "n_divergent",
    "prop_high_pareto_k",
    "min_bfmi",
    "max_bfmi",
    "total_tokens",
    "cumulative_tokens",
    "loo_se",
]

# Thresholds for conditional formatting (red if failing)
THRESHOLDS = {
    "max_r_hat":          (">=", 1.05),
    "min_ess_bulk":       ("<",  400),
    "min_ess_tail":       ("<",  100),
    "n_divergent":        (">",  0),
    "prop_high_pareto_k": (">",  0.2),
    "min_bfmi":           ("<=", 0.3),
}

# Fill colours
FILL_MEAN  = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")  # blue-grey
FILL_STD   = PatternFill(start_color="EBF1DE", end_color="EBF1DE", fill_type="solid")  # light green
FILL_META  = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")  # grey
FILL_RED   = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")  # red fail
FILL_HEAD  = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")  # dark blue
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"),  bottom=Side(style="thin"),
)


def _select_metrics(df: pd.DataFrame) -> list:
    """Return ordered list of numeric metric columns to aggregate."""
    numeric = df.select_dtypes(include=[np.number]).columns.tolist()
    meta = {"seed", "iteration"}
    available = [c for c in numeric if c not in meta]

    # Priority metrics first (those that exist), then remainder alphabetically
    ordered = [m for m in PRIORITY_METRICS if m in available]
    extras  = sorted(c for c in available if c not in ordered)
    return ordered + extras


def compute_stats(df: pd.DataFrame) -> tuple:
    """
    Group by model + dataset and compute mean, std, and count for every numeric metric.
    
    Returns:
        Tuple of (stats_df, metrics_list)
    """
    metrics = _select_metrics(df)
    rows = []

    for (model, dataset), group in df.groupby(["model", "dataset"], sort=True):
        row = {
            "model":   model,
            "dataset": dataset,
            "n_seeds": len(group),
        }
        for metric in metrics:
            if metric not in group.columns:
                row[f"{metric}_mean"] = None
                row[f"{metric}_std"]  = None
                continue
            valid = group[metric].dropna()
            row[f"{metric}_mean"] = round(valid.mean(), 4) if len(valid) > 0 else None
            row[f"{metric}_std"]  = round(valid.std(),  4) if len(valid) > 1 else 0.0
        rows.append(row)

    return pd.DataFrame(rows), metrics


def _set_header_style(ws, row_num: int, ncols: int):
    """Apply dark-blue bold white header style to a row."""
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row_num, column=c)
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = FILL_HEAD
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border    = THIN_BORDER


def _fails_threshold(metric: str, value) -> bool:
    """Return True if this mean value fails the diagnostic threshold."""
    if metric not in THRESHOLDS or value is None or (isinstance(value, float) and np.isnan(value)):
        return False
    op, threshold = THRESHOLDS[metric]
    if   op == ">=": return value >= threshold
    elif op == ">":  return value > threshold
    elif op == "<":  return value < threshold
    elif op == "<=": return value <= threshold
    return False


def write_excel(stats_df: pd.DataFrame, metrics: list, output_path: str):
    """Write both sheets to the output Excel file."""
    # Use a temp DataFrame to create the file via pandas
    stats_df.to_excel(output_path, index=False, sheet_name="_tmp")

    wb = load_workbook(output_path)
    # Remove placeholder sheet
    if "_tmp" in wb.sheetnames:
        del wb["_tmp"]

    ws_summary   = wb.create_sheet("Summary")
    ws_formatted = wb.create_sheet("Formatted (mean ± std)")

    # Write Summary sheet
    headers = ["model", "dataset"]
    for m in metrics:
        headers.append(f"{m}\nmean")
        headers.append(f"{m}\nstd")

    for c_idx, h in enumerate(headers, start=1):
        ws_summary.cell(row=1, column=c_idx).value = h
    _set_header_style(ws_summary, 1, len(headers))

    for r_idx, (_, row) in enumerate(stats_df.iterrows(), start=2):
        ws_summary.cell(row=r_idx, column=1).value = row["model"]
        ws_summary.cell(row=r_idx, column=2).value = row["dataset"]

        for col_idx in range(1, 3):
            ws_summary.cell(row=r_idx, column=col_idx).fill   = FILL_META
            ws_summary.cell(row=r_idx, column=col_idx).font   = Font(bold=True, size=10)
            ws_summary.cell(row=r_idx, column=col_idx).border = THIN_BORDER

        for m_idx, metric in enumerate(metrics):
            mean_col = 3 + m_idx * 2
            std_col = mean_col + 1
            mean_val = row.get(f"{metric}_mean")
            std_val  = row.get(f"{metric}_std")

            mean_cell = ws_summary.cell(row=r_idx, column=mean_col)
            mean_cell.value  = mean_val
            mean_cell.fill   = FILL_RED if _fails_threshold(metric, mean_val) else FILL_MEAN
            mean_cell.font   = Font(size=10)
            mean_cell.border = THIN_BORDER
            if isinstance(mean_val, float):
                mean_cell.number_format = "0.0000"

            std_cell = ws_summary.cell(row=r_idx, column=std_col)
            std_cell.value  = std_val
            std_cell.fill   = FILL_STD
            std_cell.font   = Font(size=10, color="555555")
            std_cell.border = THIN_BORDER
            if isinstance(std_val, float):
                std_cell.number_format = "0.0000"

    ws_summary.freeze_panes = "A2"
    ws_summary.column_dimensions["A"].width = 28
    ws_summary.column_dimensions["B"].width = 18
    
    # Write Formatted sheet
    headers = ["model", "dataset"] + metrics
    for c_idx, h in enumerate(headers, start=1):
        ws_formatted.cell(row=1, column=c_idx).value = h
    _set_header_style(ws_formatted, 1, len(headers))

    for r_idx, (_, row) in enumerate(stats_df.iterrows(), start=2):
        ws_formatted.cell(row=r_idx, column=1).value = row["model"]
        ws_formatted.cell(row=r_idx, column=2).value = row["dataset"]

        for col_idx in range(1, 3):
            ws_formatted.cell(row=r_idx, column=col_idx).fill   = FILL_META
            ws_formatted.cell(row=r_idx, column=col_idx).font   = Font(bold=True, size=10)
            ws_formatted.cell(row=r_idx, column=col_idx).border = THIN_BORDER

        for m_idx, metric in enumerate(metrics):
            col_idx = 3 + m_idx
            mean_val = row.get(f"{metric}_mean")
            std_val  = row.get(f"{metric}_std")

            if mean_val is None or (isinstance(mean_val, float) and np.isnan(mean_val)):
                display = "—"
            elif std_val is None or (isinstance(std_val, float) and np.isnan(std_val)) or std_val == 0:
                display = f"{mean_val:.4f}"
            else:
                display = f"{mean_val:.4f} ± {std_val:.4f}"

            cell = ws_formatted.cell(row=r_idx, column=col_idx)
            cell.value  = display
            cell.fill   = FILL_RED if _fails_threshold(metric, mean_val) else FILL_MEAN
            cell.font   = Font(size=10)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="center")

    ws_formatted.freeze_panes = "A2"
    ws_formatted.column_dimensions["A"].width = 28
    ws_formatted.column_dimensions["B"].width = 18

    wb.save(output_path)
    print(f"  Saved: {output_path}")


def main():
    parser = argparse.ArgumentParser(
        description="Generate aggregated statistics from all_seeds_best_programs_summary.csv"
    )
    parser.add_argument(
        "--input", "-i", type=str, required=True,
        help="Path to all_seeds_best_programs_summary.csv"
    )
    parser.add_argument(
        "--output", "-o", type=str, default=None,
        help="Output xlsx path (default: <input_dir>/aggregated_stats.xlsx)"
    )
    args = parser.parse_args()

    if not os.path.exists(args.input):
        print(f"Error: File not found: {args.input}")
        sys.exit(1)

    output_path = args.output or os.path.join(
        os.path.dirname(args.input), "aggregated_stats.xlsx"
    )

    # Load
    df = pd.read_csv(args.input)
    print(f"\n  Input  : {args.input}")
    print(f"  Rows   : {len(df)}")
    print(f"  Datasets: {sorted(df['dataset'].unique())}")
    print(f"  Seeds per dataset:")
    for (model, dataset), grp in df.groupby(["model", "dataset"]):
        print(f"    {dataset:<20} {len(grp)} seed(s)  {sorted(grp['seed'].tolist())}")

    # Compute stats
    stats_df, metrics = compute_stats(df)

    print(f"\n  Metrics aggregated: {len(metrics)}")
    print(f"  Groups  : {len(stats_df)}  (model × dataset combinations)")

    # Write Excel
    write_excel(stats_df, metrics, output_path)

    # Print and save quick summary
    summary_lines = []
    summary_lines.append(f"\n{'═' * 70}")
    summary_lines.append("  Aggregated Statistics (mean across seeds)")
    summary_lines.append(f"{'═' * 70}")
    display_cols = ["dataset", "reliability_score_mean", "elpd_loo_mean",
                    "max_r_hat_mean", "min_ess_bulk_mean", "n_divergent_mean"]
    available_cols = [c for c in display_cols if c in stats_df.columns]
    summary_lines.append(stats_df[available_cols].to_string(index=False))
    summary_lines.append(f"{'═' * 70}\n")
    
    summary_text = "\n".join(summary_lines)
    print(summary_text)
    
    # Save summary to file
    summary_file = os.path.join(os.path.dirname(output_path), "aggregated_summary.txt")
    with open(summary_file, "w") as f:
        f.write(summary_text)
    print(f"  Saved summary: {summary_file}")


if __name__ == "__main__":
    main()
